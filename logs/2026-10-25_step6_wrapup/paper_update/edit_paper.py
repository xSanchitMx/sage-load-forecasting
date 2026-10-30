"""Step-6 update of paper/paper.docx to the final system (MELO + drift + follow-ups + agent).
Run with system python3 (python-docx, lxml). Reads paper_numbers.json (computed from the
result files) and the new figures; writes paper/paper.docx in place (backup in _archive)."""
import copy, json, os, re, sys
from lxml import etree
import docx
from docx.shared import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from omml import omath, MNS, WNS

PROJ = "/mnt/c/Users/mukhe/Desktop/VIT Chennai/Year 4/Predictive Analytics/Project"
SRC = os.path.join(PROJ, "_archive", "paper (before step-6 update)", "paper.docx")
DST = os.path.join(PROJ, "paper", "paper.docx")
FIG = os.path.join(PROJ, "paper", "figures_v2")
N = json.load(open(os.path.join(HERE, "paper_numbers.json")))
# 2026-10-29: the final system is named SAGE (Self-Adaptive aGEntic forecaster), MAPE is the
# primary metric, the hourly-refit rows are dropped (not a feasible operating point), and the
# no-agent variant is no longer tabulated (one ablation sentence remains in Section IV-F).
FINAL = "SAGE (proposed)"
DROP = ("hourly", "Ours (no agent)", "Ours vs MELO", "Ours + agent vs ours")
for k in ("table2", "table3", "table4"):
    N[k] = [r for r in N[k] if not any(d in r[0] for d in DROP)]
    for r in N[k]:
        r[0] = {"Ours + agent": FINAL, "Ours + agent vs MELO": "SAGE vs MELO"}.get(r[0], r[0])
DS = N["det_sage"]
N["table3"][-1] = ["Caught, median delay", "–"] + [f"{DS[s]['detected']}, {DS[s]['delay_median']:.0f} h"
                                                   for s in ("level", "ramp", "amplitude", "temporary")]

W = lambda t: f"{{{WNS}}}{t}"
d = docx.Document(SRC)
P = list(d.paragraphs)                      # original indices, element refs stay valid
T2 = d.tables[1]._tbl


def text_of(p):
    return "".join(t.text or "" for t in p._p.iter(W("t")))


# ------------------------------------------------------------------ builders
def clear_content(p_el):
    for ch in list(p_el):
        if ch.tag != W("pPr"):
            p_el.remove(ch)


def wrun(text, rpr=None, bold=False, italic=False):
    r = etree.Element(W("r"))
    if rpr is not None or bold or italic:
        rp = copy.deepcopy(rpr) if rpr is not None else etree.SubElement(r, W("rPr"))
        if rpr is not None:
            r.append(rp)
        if bold and rp.find(W("b")) is None:
            etree.SubElement(rp, W("b"))
        if italic and rp.find(W("i")) is None:
            etree.SubElement(rp, W("i"))
    t = etree.SubElement(r, W("t")); t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def fill(p_el, markup, rpr=None):
    """markup: text with $latex$ inline math and **bold** spans."""
    for k, seg in enumerate(re.split(r"(\$[^$]+\$)", markup)):
        if not seg:
            continue
        if seg.startswith("$"):
            p_el.append(omath(seg[1:-1]))
        else:
            for j, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", seg)):
                if part.startswith("**"):
                    p_el.append(wrun(part[2:-2], rpr, bold=True))
                elif part:
                    p_el.append(wrun(part, rpr))


def para(template, markup, rpr=None):
    p = copy.deepcopy(template._p if hasattr(template, "_p") else template)
    clear_content(p)
    fill(p, markup, rpr)
    return p


def eq_para(latex, num):
    p = copy.deepcopy(P[79]._p); clear_content(p)
    r = etree.SubElement(p, W("r")); etree.SubElement(r, W("tab"))
    p.append(omath(latex))
    r = etree.SubElement(p, W("r")); etree.SubElement(r, W("tab"))
    t = etree.SubElement(r, W("t")); t.text = f"({num})"
    return p


def first_rpr(p_el):
    r = p_el.find(W("r"))
    return copy.deepcopy(r.find(W("rPr"))) if r is not None and r.find(W("rPr")) is not None else None


def set_text(p, markup, keep_rpr=True):
    rpr = first_rpr(p._p) if keep_rpr else None
    clear_content(p._p); fill(p._p, markup, rpr)


def replace_in_runs(p, old, new):
    for t in p._p.iter(W("t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new); return
    raise ValueError(f"not found in paragraph: {old[:60]}")


def insert_after(ref, els):
    for e in els:
        ref.addnext(e); ref = e
    return ref


def remove(*ps):
    for p in ps:
        el = p._p if hasattr(p, "_p") else p
        el.getparent().remove(el)


# ------------------------------------------------------------------ templates (copied before edits)
TPL = {k: copy.deepcopy(P[i]._p) for k, i in dict(
    h1=134, h2=136, first=141, body=142, fig=143, figcap=144, tabcap=135, alg1=97, alg2=114,
    algcap=113, ref=188).items()}
ALG_RPR = first_rpr(P[97]._p)

# ------------------------------------------------------------------ front matter
set_text(P[2], "**Abstract**—Short-term electricity load forecasting is normally performed with a single "
    "model that is trained offline and then deployed unchanged, so it degrades silently when consumption "
    "shifts. This report presents SAGE (Self-Adaptive aGEntic forecaster), a real-time, self-adaptive agentic system that removes "
    "this failure mode. "
    "A pool of five inexpensive forecasters (seasonal-naive, Holt–Winters, seasonal ARIMA, gradient-boosted "
    "trees and an online linear learner) is combined online by MELO, a multi-horizon aggregation method. An "
    "agentic layer monitors the residual of the served forecast with ADWIN and the Page–Hinkley test, refits "
    "the batch models immediately and at follow-up times after each alarm, and uses a contextual serving "
    "agent—a full-information contextual bandit—to re-weight the combiner and the individual models every "
    "hour. Evaluated prequentially on eight clients of the UCI ElectricityLoadDiagrams20112014 dataset, with "
    "every setting chosen on validation data, SAGE reaches 4.70% test MAPE, the primary metric, against 4.82% for MELO and "
    "5.01% for the initial rule-based controller, and it is significantly more accurate than MELO under "
    "Diebold–Mariano tests. Under four injected demand shifts it lowers MAPE and CV-RMSE relative to MELO in every "
    "scenario and matches daily refitting with about a fifth of the refits, at roughly 31 ms per hourly step "
    "on a laptop CPU. The contribution is an integrated, reproducible adaptive forecasting system and a "
    "controlled evaluation of each of its parts, rather than a new learning algorithm.", keep_rpr=False)
# keep the abstract paragraph's style: bold label run + plain text (style handles size/italics)
set_text(P[3], "**Index Terms**—short-term load forecasting, concept drift, online learning, forecast "
    "combination, contextual bandit, adaptive ensembles, agentic systems.", keep_rpr=False)

set_text(P[9], "An integrated self-adaptive architecture that manages a heterogeneous pool of inexpensive "
    "forecasters—rather than adapting one network—by coupling an online combiner (MELO [16]) to an explicit "
    "statistical drift detector whose alarms trigger immediate and follow-up refits.")
set_text(P[10], "A contextual serving agent, a full-information contextual bandit over the combiner and the "
    "individual models, which replaces the initial two-arm bandit controller.")
set_text(P[12], "An evaluation under a realistic streaming protocol that includes four online aggregation "
    "baselines and MELO, a daily-refitting reference, four injected demand-shift scenarios, and "
    "Diebold–Mariano significance tests.")
replace_in_runs(P[13], "setup, the dataset, and preliminary baseline results.",
                "setup, the dataset and the results, and Section V concludes.")

set_text(P[33], "Fig. 1. Left-to-right data-flow view of the proposed system. A pool of five forecasters "
    "feeds the agentic controller, where MELO combines them and a contextual agent re-weights MELO and the "
    "individual models each hour; a demand-shift detector monitors the served forecast’s error and triggers "
    "immediate and follow-up refits, and the evaluation stage returns per-candidate losses. Blue: forward "
    "forecasting path; red: drift-triggered adaptation; amber: accuracy feedback.")
replace_in_runs(P[34], "whether to keep serving its current blend of models, promote a better-performing "
                "model, or trigger a retraining event.",
                "how much to trust the combined forecast versus each individual model, and when to refit.")
replace_in_runs(P[35], "fuses the pool’s forecasts, tracks the recent accuracy of every model, and selects an "
                "action at each step—first through a deterministic rule set and then through a lightweight "
                "learned policy layered on top. Module\xa05 (Evaluation and Reward Engine) scores every served "
                "forecast, computes the accuracy, regret and adaptation-speed metrics, and returns the scalar "
                "reward that drives the learned policy.",
                "fuses the pool’s forecasts with the MELO combiner, refits the batch models after a drift alarm, "
                "and re-weights the combiner and the individual models every hour through a contextual serving "
                "agent. Module\xa05 (Evaluation and Reward Engine) scores every served forecast, computes the "
                "accuracy, regret and adaptation-speed metrics, and returns every candidate’s loss to the agent.")

# ------------------------------------------------------------------ III.C drift calibration
insert_after(P[73]._p, [para(TPL["body"],
    "Calibration. Before entering the detectors the scale-normalised residual $|e_{t}|/s_{t}$ is clipped at 10, "
    "so that a single outage hour cannot raise an alarm on its own, and the Page–Hinkley test is one-sided, "
    "because a fall in error is no reason to refit. Its threshold ($\\lambda=50$, at least 24 observations after "
    "a reset) and ADWIN’s confidence ($\\delta=0.002$) were chosen on the validation block; the first, "
    f"uncalibrated setting fired on a fixed weekly rhythm rather than on real changes. The calibrated detectors "
    f"raise about {N['false_alarms']:.1f} alarms per week on the real data.")])

# ------------------------------------------------------------------ III.D controller (rebuild 75..120)
set_text(P[75], "Module 4 is the decision-making core of SAGE, built in three layers: a MELO combiner [16] that fuses "
    "the pool, a drift response that decides when to refit, and a contextual serving agent that decides what "
    "to serve. Rolling performance tracking: the controller keeps an exponentially weighted moving average of "
    "each model’s absolute error [14],")
new = [
    para(TPL["first"], "Forecast fusion (MELO). $K=16$ correction experts each fit a time-varying affine "
         "combination of the pool forecasts with an intercept, $\\tilde{z}_{s}=(\\hat{y}_{1,s},\\dots,\\hat{y}_{M,s},1)$, "
         "by exponentially weighted least squares with forgetting factor $\\gamma_{k}$,"),
    eq_para(r"\hat{w}_{k,t}=\argmin_{w}\sum_{s<t}\gamma_{k}^{t-1-s}\left(y_{s}-w^{\top}\tilde{z}_{s}\right)^{2}", 16),
    para(TPL["body"], "Equation (16) is solved recursively for 15 forgetting factors with memory lengths "
         "$1/(1-\\gamma_{k})$ from 20 to 5,000 h plus $\\gamma_{k}=1$, and expert $k$ forecasts "
         "$\\hat{y}^{(k)}_{t}=\\hat{w}_{k,t}^{\\top}\\tilde{z}_{t}$. The intercept lets an expert absorb a level shift "
         "without retraining. The pool forecasts and the 16 experts form a candidate set $c_{j,t}$ of 21 forecasts, "
         "mixed by MLpol [28],"),
    eq_para(r"\pi_{j,t}=\frac{[R_{j,t-1}]_{+}}{\sum_{i}[R_{i,t-1}]_{+}},\quad \hat{y}^{c}_{t}=\sum_{j}\pi_{j,t}c_{j,t}", 17),
    para(TPL["first"], "where $[\\cdot]_{+}$ is the positive part and each candidate’s regret accumulates with the "
         "squared-loss gradient,"),
    eq_para(r"R_{j,t}=R_{j,t-1}+2(\hat{y}^{c}_{t}-y_{t})(\hat{y}^{c}_{t}-c_{j,t})", 18),
    para(TPL["body"], "MLpol has no learning rate. MELO’s only setting, its covariance-inflation scale, was chosen "
         "on the validation block, and the expert covariances are re-symmetrised and trace-capped for numerical "
         "stability over more than 9,000 hourly steps."),
    para(TPL["first"], "Drift response. On a drift flag the retrainable models (Holt–Winters, SARIMA, LightGBM) "
         "are refit at once on their usual windows, and follow-up refits are scheduled 24, 72 and 168 h later, "
         "so that each refit sees more of the new regime; the same models are also refit on a fixed 168-h "
         "schedule. This response was chosen on validation data with injected shifts over two alternatives that "
         "lost: refitting on post-change data only (one week is too little for SARIMA and LightGBM) and resetting "
         "MELO’s regrets and covariances after an alarm."),
    para(TPL["body"], "Contextual serving agent. On top, an agent decides each hour how much to trust the "
         "combiner versus each individual model. Its candidates are $a_{0,t}=\\hat{y}^{c}_{t}$ and "
         "$a_{m,t}=\\hat{y}_{m,t}$, each scored by a scale-free loss,"),
    eq_para(r"\ell_{j,t}=\min\left(|y_{t}-a_{j,t}|/s_{t},\,10\right)", 19),
    para(TPL["body"], "where $s_{t}$ is a running (EWMA) scale of the combiner’s absolute error. Every candidate is "
         "computed every hour, so all losses are observed after each step (full-information feedback) and no "
         "exploration is needed. Each candidate’s next loss is predicted from a context vector $x_{t}$ of 28 "
         "features—hour of day, weekend and holiday flags, candidate disagreement, fast and slow EWMA losses, "
         "deviations from the combiner, recent level change, and time since the last alarm and refit—by "
         "recursive least squares with forgetting factor $\\lambda$,"),
    eq_para(r"\theta_{j}\leftarrow\theta_{j}+g_{t}(\ell_{j,t}-\theta_{j}^{\top}x_{t}),\quad g_{t}=\frac{P x_{t}}{\lambda+x_{t}^{\top}P x_{t}}", 20),
    para(TPL["body"], "with $P\\leftarrow(P-g_{t}x_{t}^{\\top}P)/\\lambda$ shared by all candidates and its trace "
         "capped. The served forecast mixes the candidates by a softmin of predicted loss with temperature $\\tau$,"),
    eq_para(r"\omega_{j,t}=\frac{\exp(-\theta_{j}^{\top}x_{t}/\tau)}{\sum_{i}\exp(-\theta_{i}^{\top}x_{t}/\tau)},\quad \hat{y}_{t}=\sum_{j}\omega_{j,t}a_{j,t}", 21),
    para(TPL["body"], "Validation selected $\\lambda=0.999$, $\\tau=0.3$, soft mixing rather than hard selection, and "
         "absolute rather than squared loss. Because each decision affects only the current hour and its outcome "
         "is observed at once, a full-information contextual bandit is the appropriate tool rather than full "
         "reinforcement learning. The agent replaces the two-arm Discounted-UCB1 bandit of the initial design "
         "[26], [27], which only chose between the blend and the single best model. Algorithms 1 and 2 specify "
         "the control loop and the agent."),
]
ac1 = copy.deepcopy(P[96]._p); clear_content(ac1)
ac2 = copy.deepcopy(P[113]._p); clear_content(ac2)
lab_rpr = first_rpr(P[113]._p)
rest_rpr = copy.deepcopy(P[113]._p.findall(W("r"))[1].find(W("rPr"))) if P[113]._p.findall(W("r"))[1].find(W("rPr")) is not None else None
for el, label, rest in ((ac1, "Algorithm 1.", " Per-step control loop of the final system"),
                        (ac2, "Algorithm 2.", " Contextual serving agent (full-information contextual bandit)")):
    el.append(wrun(label, lab_rpr)); el.append(wrun(rest, rest_rpr))
alg1 = ["each model $m$ predicts $\\hat{y}_{m,t}$",
        "MELO forms the experts and the combined forecast $\\hat{y}^{c}_{t}$ by (16)–(17)",
        "agent builds the context $x_{t}$ and serves $\\hat{y}_{t}$ by (21)",
        "observe $y_{t}$; $e_{t}=y_{t}-\\hat{y}_{t}$",
        "update the models, MELO’s regrets (18), the agent (19)–(20) and the scores (15)",
        "feed the clipped $|e_{t}|/s_{t}$ to ADWIN and Page–Hinkley",
        "if drift flagged then",
        "refit {HW, SARIMA, LightGBM}; schedule follow-ups at $t+24$, $t+72$, $t+168$",
        "end if",
        "if a follow-up is due or 168 h have passed since the last refit then",
        "refit {HW, SARIMA, LightGBM}",
        "end if",
        "return $\\hat{y}_{t}$"]
alg2 = ["initialise $\\theta_{j}=0$ and $P=I$ for the candidates $j\\in\\{c,1,\\dots,M\\}$",
        "for each step $t$ do",
        "build $x_{t}$ (calendar, disagreement, recent losses, level change, time since alarm and refit)",
        "predict the losses $\\theta_{j}^{\\top}x_{t}$; serve the softmin mix (21)",
        "observe $y_{t}$; compute $\\ell_{j,t}$ for every candidate (19)",
        "update every $\\theta_{j}$ and $P$ by forgetting RLS (20)",
        "end for"]
block = new + [ac1] + [para(TPL["alg1"], s, ALG_RPR) for s in alg1] + [ac2] + \
        [para(TPL["alg2"], s, ALG_RPR) for s in alg2]
insert_after(P[77]._p, block)
remove(*P[78:121])

# ------------------------------------------------------------------ III.E
replace_in_runs(P[122], "returns the bandit reward.", "returns every candidate’s loss (19) to the agent.")
r131 = P[131]._p
for om in r131.findall(f"{{{MNS}}}oMath"):
    om.addprevious(omath(r"T_{\text{rec}}=\min\{\tau\le 672:\ \bar{r}_{t_{0}+\tau+u}\le 1.25\,\bar{r}_{\text{pre}},\ u=0,\dots,23\},"))
    r131.remove(om)
set_text(P[132], "Equation (26) defines the recovery time after a change at $t_{0}$: the smallest lag $\\tau$ from "
    "which the 24-h rolling mean relative error $\\bar{r}=|e_{t}|/y_{t}$ stays within 25% of its median over the "
    "two weeks before the change for 24 consecutive hours, capped at four weeks. Relative error is used so that a "
    "level shift does not count as permanent damage; events that never recover within the cap are counted "
    "separately.")
P[133]._p.addnext(para(TPL["body"],
    "Shift experiment and significance. To test adaptation under known changes, four demand shifts are "
    "injected into each client’s test series at 12 and 33 weeks: level (×1.3), ramp (to ×1.3 over four weeks), "
    "daily amplitude (×1.6) and a one-week shock (×1.4). The second event reverses the first, except for the "
    "shock, which repeats. Differences between methods are tested with the Diebold–Mariano test [30] on hourly "
    "losses (Newey–West variance, 24-h lag) for every client and scenario, and with a Wilcoxon signed-rank test "
    "across the 40 resulting series."))
remove(P[133])          # the dangling "with events that never recover..." fragment

# ------------------------------------------------------------------ Section IV
set_text(P[135], "TABLE II. Test accuracy on the real data (mean over eight clients) and model refits per week.")


def fill_table(tbl_el, header, rows, bold_rows=()):
    trs = tbl_el.findall(W("tr"))
    hdr, body_tpl = trs[0], trs[1]
    for tr in trs[1:]:
        tbl_el.remove(tr)
    for tc, val in zip(hdr.findall(W("tc")), header):
        _set_cell(tc, val)
    for i, row in enumerate(rows):
        tr = copy.deepcopy(body_tpl)
        for tc, val in zip(tr.findall(W("tc")), row):
            _set_cell(tc, val, bold=i in bold_rows)
        tbl_el.append(tr)


def _set_cell(tc, val, bold=False):
    p = tc.find(W("p")); rpr = first_rpr(p)
    for extra in tc.findall(W("p"))[1:]:
        tc.remove(extra)
    clear_content(p); p.append(wrun(val, rpr, bold=bold))


# MAE / RMSE one decimal, CV-RMSE / MAPE two decimals
t2rows = [[r[0]] + [f"{v:.2f}" if k in (2, 3) else f"{v:.1f}" for k, v in enumerate(r[1:5])] + [f"{r[5]:.1f}"]
          for r in N["table2"]]
fill_table(T2, ["Method", "MAE", "RMSE", "CV-RMSE (%)", "MAPE (%)", "Refits/wk"], t2rows,
           bold_rows=(len(t2rows) - 1,))

set_text(P[137], "All experiments are implemented in Python and run on a laptop CPU with no GPU: Holt–Winters and "
    "SARIMA use statsmodels, the gradient-boosted member uses LightGBM, and the online learner and both "
    "detectors use river. Every method is evaluated under the same prequential protocol on the test block (9,207 hours per "
    "client), and every setting—the drift thresholds, MELO’s inflation scale, the alarm response, and the "
    "agent’s $\\lambda$ and $\\tau$—was chosen on the validation block, including validation copies of the shift "
    "scenarios, before the test block was run once. The compared methods are the five single models; a "
    "validation-chosen single model; an equal-weight ensemble; four online aggregation baselines, "
    "Follow-the-Leader and Hedge [29], [31], Dynamic Weighted Majority [14] and MLpol [28]; MELO [16], replayed "
    "over the same pool forecasts with weekly refits; the initial controller (top-3 inverse-error blend with a "
    "Discounted-UCB1 bandit [26], [27]); MELO with daily refits under the shift scenarios; and the proposed "
    "system, SAGE. MAPE is the primary accuracy metric, because it is "
    "scale-free and so weights small and large clients equally; CV-RMSE is reported alongside it because MAPE is "
    "inflated in hours with near-zero load.")
set_text(P[140], "C. Accuracy on the Real Data")

# figures: new image parts
def add_fig(path, width_emu, tpl, docpr_id):
    rId, img = d.part.get_or_add_image(path)
    p = copy.deepcopy(tpl)
    h = int(width_emu * img.px_height / img.px_width)
    for blip in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
    for ext in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"):
        ext.set("cx", str(width_emu)); ext.set("cy", str(h))
    for ext in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
        if ext.get("cx"):
            ext.set("cx", str(width_emu)); ext.set("cy", str(h))
    for dp in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
        dp.set("id", str(docpr_id)); dp.set("name", f"Picture {docpr_id}")
    for c in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr"):
        c.set("name", os.path.basename(path))
    return p

COLW = 3150000
pc = N["per_client"]; pm = N["per_client_mape"]; t2 = {r[0]: r for r in N["table2"]}
rec = N["recovery"]; cs = N["cases"]
sec = [
    para(TPL["first"], f"Table II and Fig. 3 compare all methods on the real test data. SAGE has the lowest MAPE "
         f"({t2[FINAL][4]:.2f}%) and the lowest CV-RMSE ({t2[FINAL][3]:.2f}%) of every method: MELO reaches "
         f"{t2['MELO [16]'][4]:.2f}% and {t2['MELO [16]'][3]:.2f}%, the best other online aggregation method, DWM, "
         f"{t2['DWM'][4]:.2f}% MAPE, the best single model chosen on validation {t2['Validation-chosen single'][4]:.2f}%, "
         f"and the equal-weight ensemble {t2['Equal-weight ensemble'][4]:.2f}%. Its MAE is slightly above MELO’s, "
         f"because the agent leans on the individual models in low-load hours, where percentage errors are "
         f"largest. Every online aggregation method beats the initial top-3 blend, which is why MELO replaced it "
         f"as the combiner."),
    add_fig(os.path.join(FIG, "fig_accuracy.png"), COLW, TPL["fig"], 201),
    para(TPL["figcap"], "Fig. 3. Test MAPE by method (mean of eight clients); blue: SAGE, orange: MELO."),
    para(TPL["body"], f"Fig. 4 compares SAGE with the four strongest baselines client by client. SAGE has lower CV-RMSE "
         f"than Hedge, DWM and MLpol on every client and lower MAPE on seven of the eight (on MT_269 DWM is "
         f"0.001 point ahead); it beats MELO on MAPE on {N['n_agent_better_mape']} of the eight. On the four most regular clients "
         f"it is within 0.1 point of MELO, with MELO marginally ahead; its large gains come on the two irregular "
         f"clients examined in Section IV-E."),
    add_fig(os.path.join(FIG, "fig_per_client.png"), COLW, TPL["fig"], 202),
    para(TPL["figcap"], "Fig. 4. Per-client test MAPE for SAGE and the four strongest baselines: Hedge, DWM, MLpol and MELO."),
    para(TPL["h2"], "D. Adaptation Under Injected Demand Shifts"),
    para(TPL["first"], f"Table III and Fig. 5 report the shift experiment. SAGE has lower MAPE and lower CV-RMSE than "
         f"MELO in every scenario, while Hedge, DWM and MLpol are worse than MELO in every scenario. SAGE also has lower "
         f"MAPE than MELO refitted every day, in every scenario, while refitting about five times less often; on "
         f"CV-RMSE the two are level. The detector caught "
         f"{DS['level']['detected'].replace('/', ' of ')} level shifts and "
         f"{DS['temporary']['detected'].replace('/', ' of ')} temporary shocks (median delays "
         f"{DS['level']['delay_median']:.0f} h and {DS['temporary']['delay_median']:.0f} h) but only "
         f"{DS['ramp']['detected'].replace('/', ' of ')} slow ramps and "
         f"{DS['amplitude']['detected'].replace('/', ' of ')} amplitude changes, with about "
         f"{N['false_alarms_sage']:.1f} false alarms per week. After temporary shocks SAGE recovered in a median "
         f"{DS['temporary']['recovery_median']:.0f} h against {rec['temporary']['melo']:.0f} h for MELO, but after "
         f"level shifts it was slower ({DS['level']['recovery_median']:.0f} h against "
         f"{rec['level']['melo']:.0f} h). In the amplitude scenario MT_135’s near-zero hours inflate MAPE, so "
         f"CV-RMSE is the fairer measure there."),
    para(TPL["tabcap"], "TABLE III. Injected demand shifts: test MAPE / CV-RMSE (%), mean over eight clients, and "
         "shift detection by SAGE (events caught of 16, median detection delay)."),
    "TABLE3",
    add_fig(os.path.join(FIG, "fig_shifts.png"), COLW, TPL["fig"], 203),
    para(TPL["figcap"], "Fig. 5. Change in test MAPE (top) and CV-RMSE (bottom) relative to MELO for Hedge, DWM, MLpol, "
         "MELO with daily refits and SAGE; below zero is better."),
    para(TPL["body"], "Fig. 6 follows one event, chosen as the median, by MAPE change against MELO over the next two "
         "weeks, of the ten level shifts the detector caught. Load on MT_132 drops back by the 30% added earlier; "
         "the detector fires after 17 h, the models are refit then and at the follow-up times, and the agent "
         "briefly shifts weight from MELO toward SARIMA and the online model while the pool re-adapts. Across all "
         "sixteen level-shift events SAGE beats MELO over the following two weeks in eight and loses in eight; "
         "most of its gain comes from normal operation and from the cases below."),
    add_fig(os.path.join(FIG, "fig_demo.png"), COLW, TPL["fig"], 204),
    para(TPL["figcap"], "Fig. 6. Injected level shift on MT_132: load and forecasts (top), 24-h rolling absolute "
         "% error with refits ▼ (middle), agent weights (bottom). Solid line: shift; dashed: alarm."),
    para(TPL["h2"], "E. Where SAGE Helps Most"),
    para(TPL["first"], f"The largest gains over MELO come from two real clients. MT_105 (Fig. 7) is the largest "
         f"CV-RMSE gain ({cs['105']['cv_sage']:.2f}% vs {cs['105']['cv_melo']:.2f}%). After four quiet weeks at "
         f"about 100 kW, its load jumped about fivefold on 23 October 2014. SARIMA tracked the new level, and the "
         f"agent moved almost all of its weight onto SARIMA within five hours; the detector fired 7 h and 15 h "
         f"after the jump, and the refits it triggered included the new data. In the weekly-refit pool used by "
         f"MELO, SARIMA fell back to the old level nine hours after the jump, and MELO’s correction experts, "
         f"which can extrapolate beyond the pool, overshot: at that hour it forecast {cs['105']['peak_melo']:,.0f} kW against an actual "
         f"{cs['105']['peak_y']:,.0f} kW. That single hour accounts for {100 * cs['105']['peak_share']:.0f}% of "
         f"the difference in squared error on this client; without it SAGE is still better ({cs['105']['cv_sage_x']:.2f}% vs "
         f"{cs['105']['cv_melo_x']:.2f}%)."),
    add_fig(os.path.join(FIG, "fig_case_105.png"), COLW, TPL["fig"], 205),
    para(TPL["figcap"], "Fig. 7. MT_105, real data: a fivefold load jump on 23 October 2014. Load and forecasts (top) "
         "and SAGE’s agent weights (bottom, 6-h mean). Solid line: the jump; dashed: drift alarms."),
    para(TPL["body"], f"MT_135 (Fig. 8) is the largest MAPE gain ({pm['MT_135']['agent']:.2f}% vs "
         f"{pm['MT_135']['melo']:.2f}%, and 7–9% lower in every shift scenario). Its load profile is inverted: it "
         f"peaks at night at about 680 kW and falls to 150–300 kW by day. In the quietest quarter of hours MELO "
         f"over-forecasts by {cs['135']['bias_melo']:.0f} kW on average, which is a large percentage of a small "
         f"load; the agent learns that in these hours the seasonal-naive forecast (the same hour one week earlier) "
         f"is more reliable and raises its weight there from {cs['135']['w_snaive_all']:.2f} to "
         f"{cs['135']['w_snaive_low']:.2f}, and cuts the bias to "
         f"{cs['135']['bias_sage']:.0f} kW and the MAPE of those hours from {cs['135']['mape_low_melo']:.1f}% to "
         f"{cs['135']['mape_low_sage']:.1f}%. The other large gains are at the morning and evening transitions "
         f"(08:00, 19:00–21:00), where it leans on LightGBM. This client supplies most of SAGE’s mean MAPE gain over "
         f"MELO."),
    add_fig(os.path.join(FIG, "fig_case_135.png"), COLW, TPL["fig"], 206),
    para(TPL["figcap"], "Fig. 8. MT_135, real test data, by hour of day: mean load (top), MAPE of MELO and SAGE "
         "(middle), and SAGE’s mean agent weights (bottom)."),
    para(TPL["h2"], "F. Statistical Significance and Cost"),
    para(TPL["first"], "Table IV summarises the Diebold–Mariano tests of SAGE against MELO. SAGE is significantly "
         "better overall on absolute, squared and percentage error (Wilcoxon p = 0.01–0.04 across the 40 series) "
         "and significantly better in 17 series on absolute and percentage error, but significantly worse in 5–8. "
         "In an ablation without the contextual agent, the drift response alone gives the same CV-RMSE gain and "
         "about 0.1 point higher MAPE; the agent’s own contribution is not significant (p = 0.25–0.33)."),
    para(TPL["tabcap"], "TABLE IV. Diebold–Mariano tests of SAGE against MELO over 40 series (8 clients × 5 "
         "scenarios): series favouring SAGE, series significant at 5% each way, and Wilcoxon signed-rank p."),
    "TABLE4",
    para(TPL["body"], f"SAGE refits its batch models about {t2[FINAL][5]:.1f} times per week, against "
         f"{t2['MELO [16]'][5]:.1f} for MELO with weekly refits and {N['melo_daily_rpw']:.0f} with daily refits, and "
         f"runs at about {N['ms_step']:.0f} ms per hourly step on a laptop CPU, of which the agentic logic—drift "
         f"detection, MELO and the agent—is {100 * N['control_share']:.1f}%."),
    para(TPL["h2"], "G. What Did Not Work"),
    para(TPL["first"], "Several designs were tested on validation data and rejected: the initial top-3 "
         "inverse-error blend, which every online aggregation method beat; refitting on post-change data only "
         "after an alarm; resetting MELO’s regrets and covariances on alarms; and training the agent on squared "
         "rather than absolute loss. The evaluation also exposed four faults in the first implementation—a "
         "seasonal-phase error in Holt–Winters, SARIMA forecasting up to 168 h ahead without new observations, a "
         "drift detector that fired on a fixed weekly timer, and a recovery metric that always returned one "
         "hour—which inflated the earlier error figures and were corrected before every result reported here."),
    para(TPL["h1"], "V. Conclusion"),
    para(TPL["first"], "This work built SAGE, an hourly, GPU-free adaptive forecasting system for client-level "
         "electricity load: a pool of five inexpensive forecasters combined by MELO, an explicit drift detector "
         "whose alarms trigger immediate and follow-up refits, and a contextual agent that re-weights the "
         "combiner and the models each hour. On eight UCI clients it has the lowest MAPE and CV-RMSE of all "
         "methods tested and is significantly more accurate than MELO, the strongest online aggregation baseline; "
         "under injected demand shifts it beats daily refitting on MAPE and matches it on CV-RMSE at a fifth of "
         "the retraining cost. Its largest gains come from sudden real level changes and irregular load profiles; "
         "on regular clients it matches MELO. Richer context such as weather and calendar events, and giving the "
         "agent control over refit timing, are natural next steps."),
]
# tables III / IV from the (already refilled) Table II as the template
T3 = copy.deepcopy(T2); fill_table(T3, ["Method", "No shift", "Level", "Ramp", "Amplitude", "Temporary"],
                                   N["table3"], bold_rows=(2,))  # row 2 = SAGE (proposed)
T4 = copy.deepcopy(T2); fill_table(T4, ["Comparison", "Loss", "Better in", "Sig. better", "Sig. worse", "Wilcoxon p"],
                                   N["table4"])
sec = [T3 if s == "TABLE3" else T4 if s == "TABLE4" else s for s in sec]
insert_after(P[140]._p, sec)
old_ids = []
for p in P[141:161]:
    for blip in p._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        old_ids.append(blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"))
remove(*P[141:161])

# Fig. 1: swap the architecture image
arch_rid = P[32]._p.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip").get(
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
arch_part = d.part.related_parts[arch_rid]
arch_part._blob = open(os.path.join(PROJ, "paper", "architecture_pictorial_v2.png"), "rb").read()

# references
refs = ["[28] P. Gaillard, G. Stoltz, and T. van Erven, “A Second-order Bound with Excess Losses,” in Proc. "
        "Conf. on Learning Theory (COLT), 2014.",
        "[29] N. Cesa-Bianchi and G. Lugosi, Prediction, Learning, and Games. Cambridge University Press, 2006.",
        "[30] F. X. Diebold and R. S. Mariano, “Comparing Predictive Accuracy,” Journal of Business & Economic "
        "Statistics, vol. 13, no. 3, 1995.",
        "[31] N. Littlestone and M. K. Warmuth, “The Weighted Majority Algorithm,” Information and "
        "Computation, vol. 108, no. 2, 1994."]
_last_ref = insert_after(P[188]._p, [para(TPL["ref"], r, first_rpr(TPL["ref"])) for r in refs])

# submission requirement (2026-10-30): dataset and code links after the References, as live hyperlinks
from docx.opc.constants import RELATIONSHIP_TYPE as _RT
DATASET_URL = "https://archive.ics.uci.edu/dataset/321/electricityloaddiagrams20112014"
GITHUB_URL = "https://github.com/xSanchitMx/sage-load-forecasting"


def link_para(label, url):
    """Reference-sized, left-aligned paragraph: bold label, line break, clickable URL."""
    p = para(TPL["ref"], "")
    ppr = p.find(W("pPr"))
    for tag in ("spacing", "ind", "jc"):
        for old in ppr.findall(W(tag)):
            ppr.remove(old)
    sp = etree.Element(W("spacing")); sp.set(W("before"), "80")
    ind = etree.Element(W("ind")); ind.set(W("left"), "0"); ind.set(W("hanging"), "0"); ind.set(W("firstLine"), "0")
    jc = etree.Element(W("jc")); jc.set(W("val"), "left")
    # schema order inside pPr: ... spacing, ind, ..., jc, ..., rPr, sectPr
    later = next((ch for ch in ppr if ch.tag in (W("textDirection"), W("textAlignment"), W("outlineLvl"),
                                                 W("rPr"), W("sectPr"), W("pPrChange"))), None)
    for el in (sp, ind, jc):
        (later.addprevious(el) if later is not None else ppr.append(el))
    rpr = first_rpr(TPL["ref"])
    p.append(wrun(label, rpr, bold=True))
    br = etree.SubElement(p, W("r")); etree.SubElement(br, W("br"))
    h = etree.SubElement(p, W("hyperlink"))
    h.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
          d.part.relate_to(url, _RT.HYPERLINK, is_external=True))
    r = etree.SubElement(h, W("r"))
    rp = copy.deepcopy(rpr) if rpr is not None else etree.Element(W("rPr"))
    r.append(rp)
    c = etree.SubElement(rp, W("color")); c.set(W("val"), "0563C1")
    u = etree.SubElement(rp, W("u")); u.set(W("val"), "single")
    t = etree.SubElement(r, W("t")); t.text = url
    return p


insert_after(_last_ref, [link_para("Dataset Link:", DATASET_URL), link_para("GitHub Link:", GITHUB_URL)])

# Fig. 2: layered pictorial architecture, right after Fig. 1's caption (added 2026-10-30)
insert_after(P[36]._p, [   # after the notation paragraph, so text fills the page under Fig. 1 first
    add_fig(os.path.join(PROJ, "paper", "architecture_icons.png"), COLW, TPL["fig"], 207),
    para(TPL["figcap"], "Fig. 2. Pictorial view of SAGE as five layers: hourly load from the customers' meters (1) "
         "feeds five forecasters (2), MELO combines them (3), the contextual agent mixes MELO and the five models "
         "using 28 context features (4; bars show its mean weights on MT_132), and the served forecast’s error drives "
         "the drift watchdog, whose alarms retrain the batch models now and 1, 3 and 7 days later (5)."),
])
insert_after(P[35]._p, [para(TPL["body"], "Fig. 1 traces the data flow between the modules; Fig. 2 shows the same "
    "system as layers, with the icon for each component and the retraining loop that returns from the watchdog to "
    "the model pool.")])

# P[35] closed a 2-column section that is immediately followed by another 2-column section.
# Word balances the columns at every continuous section end, so with Fig. 2 inside that short
# section a column-height gap appeared beside the figure. Merging the two sections removes it.
_sp = P[35]._p.find(f"{W('pPr')}/{W('sectPr')}")
if _sp is not None:
    _sp.getparent().remove(_sp)

# keep every figure on the same column as its caption, and every caption in one piece
_DRAW = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
for _p in d.element.body.iter(W("p")):
    _ppr = _p.find(W("pPr"))
    _sty = _ppr.find(W("pStyle")).get(W("val")) if _ppr is not None and _ppr.find(W("pStyle")) is not None else ""
    _flag = "keepNext" if next(_p.iter(_DRAW), None) is not None else ("keepLines" if _sty == "ImageCaption" else None)
    if _flag:
        if _ppr is None:
            _ppr = etree.SubElement(_p, W("pPr")); _p.insert(0, _ppr)
        if _ppr.find(W(_flag)) is None:
            _el = etree.Element(W(_flag))
            _after = _ppr.find(W("pStyle"))          # schema order: pStyle, keepNext, keepLines, ...
            if _flag == "keepLines" and _ppr.find(W("keepNext")) is not None:
                _after = _ppr.find(W("keepNext"))
            (_after.addnext(_el) if _after is not None else _ppr.insert(0, _el))

# Table I: start on a new page with its caption instead of leaving caption + header row at the foot of page 1
for _p in d.element.body.iter(W("p")):
    if "".join(t.text or "" for t in _p.iter(W("t"))).startswith("TABLE I.\u2002") or \
       "".join(t.text or "" for t in _p.iter(W("t"))).startswith("TABLE I."):
        _ppr = _p.find(W("pPr"))
        if _ppr.find(W("pageBreakBefore")) is None:
            _el = etree.Element(W("pageBreakBefore"))
            _anchor = _ppr.find(W("keepLines")) or _ppr.find(W("keepNext")) or _ppr.find(W("pStyle"))
            (_anchor.addnext(_el) if _anchor is not None else _ppr.insert(0, _el))
        break

for rid in set(old_ids):
    try:
        d.part.drop_rel(rid)
    except Exception as e:
        print("drop_rel", rid, e)
d.save(DST)
print("saved", DST)
